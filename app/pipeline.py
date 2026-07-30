"""Document processing pipeline: parsing, chunking, and embedding.

Provides overlapping sliding window chunking with paragraph tracking,
multi-format parsing, and integration with the TF-IDF embedding system.
"""

from __future__ import annotations

import csv
import uuid
from collections.abc import Iterable
from datetime import datetime
from pathlib import Path

from app.embeddings import embed
from app.models import Chunk

# Chunking parameters
CHUNK_SIZE = 512  # characters
CHUNK_OVERLAP = 100  # characters of overlap between chunks


def parse_content(path: str, mime_type: str) -> tuple[str, dict]:
    """Parse file content based on extension and mime type.

    Handles text files, CSVs, PDFs, DOCX, XLSX, and images.
    Returns (text_content, parse_details).
    """
    p = Path(path)

    if not p.exists():
        return "", {"parser": "missing-file", "pages": 0}

    ext = p.suffix.lower()

    # Check for binary files that might slip through
    if ext in {".bin", ".exe", ".dll", ".so", ".dylib", ".o", ".pyc", ".class"}:
        return "", {"parser": "binary-skipped", "pages": 0}

    if ext in {".txt", ".md", ".csv"} or mime_type.startswith("text/"):
        try:
            text = p.read_text(encoding="utf-8", errors="ignore")
        except (OSError, UnicodeDecodeError):
            return "", {"parser": "read-error", "pages": 0}

        if not text.strip():
            return "", {"parser": "empty-file", "pages": 0}

        details: dict = {"parser": "text", "pages": 1}
        if ext == ".csv":
            try:
                with p.open("r", encoding="utf-8", errors="ignore") as f:
                    rows = list(csv.reader(f))
                details["rows"] = max(len(rows) - 1, 0)
            except csv.Error:
                details["rows"] = 0
        return text, details

    if ext == ".pdf":
        try:
            from pypdf import PdfReader  # type: ignore

            reader = PdfReader(path)
            pages = [page.extract_text() or "" for page in reader.pages]
            text = "\n\n".join(pages)
            if not text.strip():
                return "", {"parser": "pypdf-empty", "pages": len(pages)}
            return text, {"parser": "pypdf", "pages": len(pages)}
        except Exception:
            try:
                text = p.read_text(encoding="utf-8", errors="ignore")
                if text.strip():
                    return text, {"parser": "fallback", "pages": 1}
            except (OSError, UnicodeDecodeError):
                pass
            return "", {"parser": "pdf-error", "pages": 0}

    if ext == ".docx":
        try:
            from docx import Document  # type: ignore

            doc = Document(path)
            paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]
            text = "\n\n".join(paragraphs)
            if not text.strip():
                return "", {"parser": "docx-empty", "paragraphs": 0}
            return text, {"parser": "python-docx", "paragraphs": len(paragraphs)}
        except Exception:
            return "", {"parser": "docx-unavailable", "paragraphs": 0}

    if ext in {".xlsx", ".xls"}:
        try:
            import openpyxl  # type: ignore

            wb = openpyxl.load_workbook(path)
            texts = []
            total_rows = 0
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    vals = [str(c) for c in row if c is not None]
                    if vals:
                        texts.append(" | ".join(vals))
                        total_rows += 1
            text = "\n".join(texts)
            if not text.strip():
                return "", {"parser": "xlsx-empty", "rows": 0}
            return text, {"parser": "openpyxl", "rows": total_rows}
        except Exception:
            return "", {"parser": "xlsx-unavailable", "rows": 0}

    if ext in {".png", ".jpg", ".jpeg", ".heic"}:
        return "[Image content placeholder: OCR/VQA pipeline required in production]", {
            "parser": "image-placeholder",
            "ocr": False,
        }

    # Fallback: try reading as text
    try:
        text = p.read_text(encoding="utf-8", errors="ignore")
        if not text.strip():
            return "", {"parser": "empty-file", "pages": 0}
        return text, {"parser": "fallback", "pages": 1}
    except (OSError, UnicodeDecodeError):
        return "", {"parser": "read-error", "pages": 0}


def _sliding_window_chunks(
    text: str, chunk_size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP
) -> Iterable[tuple[str, int, int]]:
    """Generate overlapping chunks using a sliding window.

    Yields (chunk_text, start_offset, end_offset) tuples.
    """
    if not text:
        return

    text_len = len(text)

    if text_len <= chunk_size:
        yield text, 0, text_len
        return

    start = 0
    while start < text_len:
        end = min(start + chunk_size, text_len)
        chunk_text = text[start:end]

        # Try to break at a word boundary if we're not at the end
        if end < text_len:
            # Look for the last space within the chunk
            last_space = chunk_text.rfind(" ")
            if last_space > chunk_size // 2:  # Only break if we keep at least half the chunk
                end = start + last_space
                chunk_text = text[start:end]

        if chunk_text.strip():
            yield chunk_text.strip(), start, end

        # Advance by chunk_size - overlap
        start = end - overlap if end < text_len else text_len


def _determine_paragraph_index(text: str, start_offset: int) -> int:
    """Determine which paragraph a given offset falls into."""
    # Count double-newline separated blocks before the offset
    prefix = text[:start_offset]
    return prefix.count("\n\n")


def chunk_document(
    file_id: str, filename: str, upload_ts: datetime, text: str, page: int = 1
) -> list[Chunk]:
    """Split document text into overlapping chunks with embeddings.

    Uses a sliding window approach (512 chars, 100 char overlap) for better
    context preservation at chunk boundaries. Each chunk tracks its paragraph
    index for citation mapping.

    Args:
        file_id: The file record ID.
        filename: Original filename.
        upload_ts: Upload timestamp.
        text: Full document text.
        page: Page number (for multi-page documents).

    Returns:
        List of Chunk objects with embeddings and metadata.
    """
    if not text or not text.strip():
        return []

    chunks: list[Chunk] = []
    for chunk_text, start, end in _sliding_window_chunks(text):
        para_idx = _determine_paragraph_index(text, start)
        chunks.append(
            Chunk(
                id=str(uuid.uuid4()),
                file_id=file_id,
                text=chunk_text,
                start_offset=start,
                end_offset=end,
                page=page,
                paragraph_index=para_idx,
                embedding=embed(chunk_text),
                metadata={
                    "filename": filename,
                    "page": page,
                    "paragraphIndex": para_idx,
                    "charOffsets": [start, end],
                    "uploadTimestamp": upload_ts.isoformat() + "Z",
                },
            )
        )
    return chunks
